"""
Pipeline principal: cruza el listado oficial de estudiantes con los
resultados crudos de ZipGrade (sesion 1 y 2, los 6 grados), calcula
puntajes por area/sesion usando answer_keys.json, y escribe los JSON
que consume el sitio estatico (data/estudiantes.json, data/promedios.json,
data/meta.json).

Solo escribe a la salida: nombre, matricula, grado, grupo y resultados
calculados. Telefono, direccion y numero de documento del listado NUNCA
se copian a la salida.

Uso:
    python3 build/extract_answer_keys.py   # primero, si answer_keys.json no existe o cambio la fuente
    python3 build/build_data.py
"""
import csv
import io
import json
import re
import unicodedata
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl

BASE = Path("/home/froylan/Documents/Pruebas periodicas CESUM/Tercer periodo 2026")
ROSTER_XLSX = BASE / "LISTADO DE 6º A 11º.xlsx"
ZIP_S1 = BASE / "Resultados_sesion_1.zip"
ZIP_S2 = BASE / "Resultados_sesion2.zip"

# Para (grado, sesion) que aparezcan aqui, se recalifica desde la respuesta
# letra-por-letra del estudiante (comparada contra answer_keys.json) en vez
# de confiar en "Points Earned" de ZipGrade. Necesario cuando la clave que
# se subio originalmente a ZipGrade para calificar tenia una desalineacion
# de numeracion (ver DROP_DUPLICATE_ROWS en extract_answer_keys.py) -- los
# puntos ya calculados por ZipGrade quedaron mal para las preguntas
# posteriores al desfase, asi que hay que recalificar con la clave corregida.
RESP_LETRA_FILES = {
    ("9", "session2"): BASE / "NovenoIIIP-S2-all-Nombre-ResptLetra-2026-09-08 19_12_34.csv",
    ("9", "session1"): BASE / "Noveno_sesion_1-all-Quiz Format 2026-09-08 01_12-2026-09-08 18_14_49 - Noveno_sesion_1-all-Quiz Format 2026-09-08 01_12-2026-09-08 18_14_49.csv",
    ("10", "session1"): BASE / "Decimo_sesion_1-all-Quiz Format 2026-09-08 01_12-2026-09-08 19_48_26.csv",
    # grado 10 sesion 2: la clave subida a ZipGrade tenia mal ED. FISICA
    # (16-21), asi que "Points Earned" quedo mal para esa area -- se recalifica
    # desde la respuesta letra-por-letra contra la clave corregida
    # (ver FINAL_KEY_OVERRIDES en extract_answer_keys.py).
    ("10", "session2"): BASE / "DecimoIIIP-S2-all-Nombre-ResptLetra-2026-09-10 19_07_17.csv",
    ("11", "session1"): BASE / "Once_sesion_1-all-Quiz Format 2026-09-08 01_12-2026-09-08 19_48_52.csv",
    ("11", "session2"): BASE / "UndecimoIIIP-S2-all-Nombre-ResptLetra-2026-09-08 20_58_14.csv",
}

BUILD_DIR = Path(__file__).parent
ANSWER_KEYS_PATH = BUILD_DIR / "answer_keys.json"
OVERRIDES_PATH = BUILD_DIR / "overrides.json"
DATA_DIR = BUILD_DIR.parent / "data"
UNMATCHED_REPORT = BUILD_DIR / "unmatched_report.txt"

GRADE_NAME_MAP = {
    "SEXTO": "6",
    "SEPTIMO": "7",
    "OCTAVO": "8",
    "NOVENO": "9",
    "DECIMO": "10",
    "ONCE": "11",
    "UNDECIMO": "11",
}

# prefijos (en minuscula) de los nombres de archivo CSV de ZipGrade por grado,
# para cada sesion (no dependemos de las marcas de tiempo del nombre completo).
ZIPGRADE_PREFIXES = {
    "session1": {
        "6": "sexto",
        "7": "septimo",
        "8": "octavo",
        "9": "noveno",
        "10": "decimo",
        "11": "once",
    },
    "session2": {
        "6": "sexto",
        "7": "septimo",
        "8": "octavo",
        "9": "noveno",
        "10": "decimo",
        "11": "undecimo",
    },
}


def strip_accents(s):
    return "".join(
        c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn"
    )


def norm(s):
    if s is None:
        return ""
    s = strip_accents(str(s)).upper()
    s = re.sub(r"\s+", " ", s).strip()
    return s


def norm_grade_label(s):
    s = strip_accents(str(s or "")).upper().strip()
    return GRADE_NAME_MAP.get(s)


# ---------------------------------------------------------------------------
# 1. Listado oficial (roster)
# ---------------------------------------------------------------------------
def load_roster():
    """Devuelve:
    roster_by_key: {(grado, sid): {"nombre":..., "grado":..., "grupo":..., "sid":...}}
    roster_by_norm_name: {grado: {nombre_normalizado: record}}
    grupos_por_grado: {grado: [lista de grupos ordenada]}

    En el listado oficial hay matriculas repetidas (dos estudiantes distintos
    con la misma matricula). Antes eso colapsaba a ambos en un solo registro y
    hacia desaparecer a uno del panel. Ahora, cuando una matricula se repite
    dentro de un mismo grado, se genera un `sid` unico
    (matricula-grupo-primerapellido) para cada estudiante; el resto conserva
    `sid == matricula`.
    """
    wb = openpyxl.load_workbook(ROSTER_XLSX, data_only=True)
    grupos_por_grado = {g: set() for g in set(GRADE_NAME_MAP.values())}

    filas = []  # (grado, grupo_id, apellidos_nombres, matricula)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        grado_label = ws["H3"].value
        grado = norm_grade_label(grado_label)
        if grado is None:
            print(f"AVISO: hoja {sheet_name} sin grado reconocible ({grado_label!r}), se omite")
            continue

        grupo_cell = ws["K1"].value or ""
        m = re.search(r"Grupo:\s*(\S+)", str(grupo_cell))
        if not m:
            print(f"AVISO: hoja {sheet_name} sin 'Grupo:' reconocible, se omite")
            continue
        letra = m.group(1).strip()[-1].upper()
        grupo_id = f"{grado}{letra}"
        grupos_por_grado[grado].add(grupo_id)

        r = 5
        while True:
            apellidos_nombres = ws.cell(row=r, column=2).value
            if apellidos_nombres is None or str(apellidos_nombres).strip() == "":
                break
            matricula = ws.cell(row=r, column=5).value
            matricula = str(matricula).strip() if matricula is not None else None
            filas.append((grado, grupo_id, str(apellidos_nombres).strip(), matricula))
            r += 1

    # matriculas repetidas dentro de un mismo grado -> hay que desambiguar
    conteo = {}
    for grado, _, _, matricula in filas:
        if matricula:
            conteo[(grado, matricula)] = conteo.get((grado, matricula), 0) + 1
    duplicadas = {k for k, n in conteo.items() if n > 1}

    roster_by_key = {}
    roster_by_norm_name = {g: {} for g in GRADE_NAME_MAP.values()}
    for grado, grupo_id, apellidos_nombres, matricula in filas:
        nombre_fmt = apellidos_nombres.title()
        if matricula and (grado, matricula) in duplicadas:
            primer_apellido = norm(apellidos_nombres).split(" ")[0] or "X"
            sid = f"{matricula}-{grupo_id}-{primer_apellido}"
            print(
                f"AVISO: matricula {matricula} repetida en grado {grado}; "
                f"{apellidos_nombres!r} recibe id sintetico {sid!r}"
            )
        else:
            sid = matricula

        record = {"nombre": nombre_fmt, "grado": grado, "grupo": grupo_id, "sid": sid}
        if sid:
            roster_by_key[(grado, sid)] = record
        roster_by_norm_name[grado][norm(apellidos_nombres)] = {
            **record,
            "matricula": matricula,
        }

    grupos_por_grado = {
        g: sorted(grupos_por_grado[g]) for g in sorted(grupos_por_grado, key=int)
    }
    return roster_by_key, roster_by_norm_name, grupos_por_grado


# ---------------------------------------------------------------------------
# 2. Overrides manuales
# ---------------------------------------------------------------------------
def load_overrides():
    raw = json.loads(OVERRIDES_PATH.read_text(encoding="utf-8"))
    by_id = {}
    for entry in raw["by_id"]:
        key = (entry["grado"], entry["session"], str(entry["zipgrade_id"]))
        by_id[key] = entry.get("sid") or entry["matricula"]
    return by_id


# ---------------------------------------------------------------------------
# 3. Resultados crudos de ZipGrade
# ---------------------------------------------------------------------------
def find_zip_member(zf, grade, session):
    prefix = ZIPGRADE_PREFIXES[session][grade]
    for name in zf.namelist():
        base = name.rsplit("/", 1)[-1]
        if base.lower().startswith(prefix.lower()) and base.lower().endswith(".csv"):
            return name
    raise FileNotFoundError(f"No se encontro CSV de ZipGrade para grado {grade} / {session}")


def load_zipgrade_csv(zip_path, grade, session):
    with zipfile.ZipFile(zip_path) as zf:
        member = find_zip_member(zf, grade, session)
        raw = zf.read(member).decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(raw)))


def load_response_letras_csv(path):
    with open(path, encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


# ---------------------------------------------------------------------------
# 4. Cruce + calificacion
# ---------------------------------------------------------------------------
def score_student_session(row, areas, n_questions, key, from_letters=False):
    """A partir de una fila de ZipGrade (dict) y la lista de areas
    [{area,start,end}], calcula el detalle por area.

    Si from_letters=True, la fila trae "#N Student Response" (la letra que
    marco el estudiante) y se recalifica comparando contra `key`, en vez de
    confiar en "#N Points Earned" (util cuando la clave subida a ZipGrade
    tenia una desalineacion de numeracion -- ver RESP_LETRA_FILES)."""
    points = {}
    if from_letters:
        for q in range(1, n_questions + 1):
            resp = (row.get(f"#{q} Student Response") or "").strip().upper()
            correcta = key.get(str(q))
            points[q] = 1.0 if (resp and correcta and resp == correcta) else 0.0
    else:
        for q in range(1, n_questions + 1):
            col = f"#{q} Points Earned"
            val = row.get(col)
            try:
                points[q] = float(val) if val not in (None, "") else 0.0
            except ValueError:
                points[q] = 0.0

    areas_out = []
    total_correct = 0.0
    total_preguntas_validas = 0
    for a in areas:
        # las preguntas sin clave (VOID_QUESTIONS: existen como columna pero
        # no tienen opciones reales de respuesta) se excluyen del area
        qs = [q for q in range(a["start"], a["end"] + 1) if key.get(str(q)) is not None]
        correct = sum(points[q] for q in qs)
        total = len(qs)
        total_correct += correct
        total_preguntas_validas += total
        preguntas = [
            {"n": q, "correcta": points[q] >= 0.999, "clave": key.get(str(q))}
            for q in qs
        ]
        areas_out.append(
            {
                "area": a["area"],
                "correctas": correct,
                "total": total,
                "porcentaje": round(correct / total * 100, 1) if total else 0.0,
                "preguntas": preguntas,
            }
        )
    pct_total = round(total_correct / total_preguntas_validas * 100, 1) if total_preguntas_validas else 0.0
    return areas_out, pct_total


def main():
    answer_keys = json.loads(ANSWER_KEYS_PATH.read_text(encoding="utf-8"))
    overrides_by_id = load_overrides()
    roster_by_key, roster_by_norm_name, grupos_por_grado = load_roster()

    estudiantes = {}  # (grado, matricula) -> record
    unmatched_lines = []
    stats = {}

    for grade in answer_keys.keys():
        for session, zip_path in (("session1", ZIP_S1), ("session2", ZIP_S2)):
            areas = answer_keys[grade][session]["areas"]
            n_q = answer_keys[grade][session]["total"]
            key = answer_keys[grade][session]["key"]

            resp_letra_path = RESP_LETRA_FILES.get((grade, session))
            from_letters = resp_letra_path is not None
            if from_letters:
                rows = load_response_letras_csv(resp_letra_path)
                print(f"  (grado {grade} {session}: recalificando desde respuesta-letra {resp_letra_path.name})")
            else:
                rows = load_zipgrade_csv(zip_path, grade, session)

            matched = 0
            unmatched = 0
            for row in rows:
                first = (row.get("Student First Name") or "").strip()
                last = (row.get("Student Last Name") or "").strip()
                zid = (row.get("Student ID") or "").strip()

                override_key = overrides_by_id.get((grade, session, zid))
                roster_rec = None
                if override_key:
                    # el override apunta directo al sid del listado (matricula o
                    # id sintetico si la matricula estaba repetida)
                    roster_rec = roster_by_key.get((grade, override_key))
                elif first or last:
                    name_key = norm(f"{last} {first}")
                    roster_rec = roster_by_norm_name[grade].get(name_key)

                if not roster_rec:
                    unmatched += 1
                    unmatched_lines.append(
                        f"grado={grade} session={session} zid={zid!r} "
                        f"nombre={first!r} {last!r} -> SIN COINCIDENCIA"
                    )
                    continue

                matched += 1
                areas_out, pct_total = score_student_session(row, areas, n_q, key, from_letters)

                sid = roster_rec.get("sid") or f"{grade}-{zid}"
                student_key = (grade, sid)
                if student_key not in estudiantes:
                    estudiantes[student_key] = {
                        "id": sid,
                        "nombre": roster_rec["nombre"],
                        "grado": grade,
                        "grupo": roster_rec["grupo"],
                        "sesiones": {},
                    }
                estudiantes[student_key]["sesiones"][session] = {
                    "areas": areas_out,
                    "porcentaje_total": pct_total,
                }

            stats[(grade, session)] = {
                "total_filas": len(rows),
                "matched": matched,
                "unmatched": unmatched,
            }
            print(
                f"grado {grade:>2} {session}: {matched}/{len(rows)} unidos "
                f"({unmatched} sin coincidencia)"
            )

    # overall pct combinando ambas sesiones si existen
    for rec in estudiantes.values():
        s1 = rec["sesiones"].get("session1")
        s2 = rec["sesiones"].get("session2")
        if s1 and s2:
            c1 = sum(a["correctas"] for a in s1["areas"])
            t1 = sum(a["total"] for a in s1["areas"])
            c2 = sum(a["correctas"] for a in s2["areas"])
            t2 = sum(a["total"] for a in s2["areas"])
            rec["porcentaje_general"] = round((c1 + c2) / (t1 + t2) * 100, 1) if (t1 + t2) else 0.0
        elif s1:
            rec["porcentaje_general"] = s1["porcentaje_total"]
        elif s2:
            rec["porcentaje_general"] = s2["porcentaje_total"]
        else:
            rec["porcentaje_general"] = None

    UNMATCHED_REPORT.write_text("\n".join(unmatched_lines) + "\n", encoding="utf-8")
    print(f"\nReporte de no coincidencias: {UNMATCHED_REPORT} ({len(unmatched_lines)} filas)")

    # ---- posiciones (ranking) por grupo y por grado, segun porcentaje_general ----
    def asignar_posiciones(recs, campo):
        ordenados = sorted(recs, key=lambda r: r["porcentaje_general"], reverse=True)
        total = len(ordenados)
        for i, r in enumerate(ordenados, start=1):
            r.setdefault("posicion", {})[campo] = {"puesto": i, "de": total}

    por_grupo_recs = {}
    por_grado_recs = {}
    for rec in estudiantes.values():
        if rec["porcentaje_general"] is None:
            continue
        por_grupo_recs.setdefault(rec["grupo"], []).append(rec)
        por_grado_recs.setdefault(rec["grado"], []).append(rec)
    for recs in por_grupo_recs.values():
        asignar_posiciones(recs, "grupo")
    for recs in por_grado_recs.values():
        asignar_posiciones(recs, "grado")

    # ---- agregados ----
    by_grado_area = {}  # grado -> session -> area -> [pcts]
    by_grupo_area = {}  # grupo -> session -> area -> [pcts]
    by_grado = {}  # grado -> [pcts generales]
    by_grupo = {}  # grupo -> [pcts generales]
    colegio_pcts = []
    colegio_por_sesion = {"session1": [], "session2": []}

    for rec in estudiantes.values():
        g = rec["grado"]
        gr = rec["grupo"]
        if rec["porcentaje_general"] is not None:
            by_grado.setdefault(g, []).append(rec["porcentaje_general"])
            by_grupo.setdefault(gr, []).append(rec["porcentaje_general"])
            colegio_pcts.append(rec["porcentaje_general"])
        for session, sdata in rec["sesiones"].items():
            colegio_por_sesion[session].append(sdata["porcentaje_total"])
            area_map_grado = by_grado_area.setdefault(g, {}).setdefault(session, {})
            area_map_grupo = by_grupo_area.setdefault(gr, {}).setdefault(session, {})
            for a in sdata["areas"]:
                area_map_grado.setdefault(a["area"], []).append(a["porcentaje"])
                area_map_grupo.setdefault(a["area"], []).append(a["porcentaje"])

    def avg(lst):
        return round(sum(lst) / len(lst), 1) if lst else None

    promedios = {
        "colegio": {
            "promedio_general": avg(colegio_pcts),
            "por_sesion": {k: avg(v) for k, v in colegio_por_sesion.items()},
            "n_estudiantes": len(estudiantes),
        },
        "por_grado": {g: avg(v) for g, v in by_grado.items()},
        "por_grupo": {g: avg(v) for g, v in by_grupo.items()},
        "por_area": {
            g: {s: {a: avg(v) for a, v in amap.items()} for s, amap in sessions.items()}
            for g, sessions in by_grado_area.items()
        },
        "por_grupo_area": {
            gr: {s: {a: avg(v) for a, v in amap.items()} for s, amap in sessions.items()}
            for gr, sessions in by_grupo_area.items()
        },
    }

    # indice global de areas: nombre de area -> lista de (grado, sesion) donde aparece
    areas_globales = {}
    for g, sess in promedios["por_area"].items():
        for session, amap in sess.items():
            for area in amap:
                areas_globales.setdefault(area, []).append({"grado": g, "session": session})
    for area in areas_globales:
        areas_globales[area].sort(key=lambda x: int(x["grado"]))

    meta = {
        "generado": datetime.now().isoformat(timespec="seconds"),
        "grados": sorted(answer_keys.keys(), key=int),
        "grupos_por_grado": grupos_por_grado,
        "areas_por_grado": {
            g: {
                "session1": [a["area"] for a in answer_keys[g]["session1"]["areas"]],
                "session2": [a["area"] for a in answer_keys[g]["session2"]["areas"]],
            }
            for g in answer_keys
        },
        "areas_globales": areas_globales,
    }

    DATA_DIR.mkdir(exist_ok=True)
    (DATA_DIR / "estudiantes.json").write_text(
        json.dumps(list(estudiantes.values()), ensure_ascii=False), encoding="utf-8"
    )
    (DATA_DIR / "promedios.json").write_text(
        json.dumps(promedios, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (DATA_DIR / "meta.json").write_text(
        json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"\nEscritos: {DATA_DIR}/estudiantes.json, promedios.json, meta.json")
    print(f"Total estudiantes con al menos 1 sesion: {len(estudiantes)}")


if __name__ == "__main__":
    main()
