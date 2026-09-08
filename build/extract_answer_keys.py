"""
Extrae las claves de respuesta (área -> rango de preguntas -> letra correcta)
de los xlsx dentro de "Pruebas completas.zip", para sesión 1 (Cuadernillo 1)
y sesión 2 (Cuadernillo 2), los 6 grados.

No contiene datos de estudiantes -- solo áreas/preguntas/claves, así que el
resultado (answer_keys.json) es seguro de commitear a un repo público.
"""
import io
import json
import zipfile
from pathlib import Path

import openpyxl

SOURCE_ZIP = Path(
    "/home/froylan/Documents/Pruebas periodicas CESUM/Tercer periodo 2026/Pruebas completas.zip"
)
OUT_PATH = Path(__file__).parent / "answer_keys.json"

GRADE_FILES = {
    "6": "III PERIODO/GRADO 6°/TABLA DE RESPUESTAS 1_P3_2026 6°.xlsx",
    "7": "III PERIODO/GRADO 7°/TABLA DE RESPUESTAS 1_P3_2026 7°.xlsx",
    "8": "III PERIODO/GRADO 8/TABLA DE RESPUESTAS CUADERNILLOS 8_P3_2026.xlsx",
    "9": "III PERIODO/GRADO 9/TABLA DE RESPUESTAS DEL CUADERNILLO_grado_9_2026.xlsx",
    "10": "III PERIODO/Grado 10/TABLA DE RESPUESTAS _Grado 10_P3_2026.xlsx",
    "11": "III PERIODO/grado 11/TABLA DE RESPUESTAS _Grado 11_P3_2026.xlsx",
}

# Claves confirmadas manualmente con el usuario cuando la celda venía vacía
# en el archivo fuente. clave: (grado, sesion, pregunta) -> letra
MANUAL_FIXES = {
    ("6", "session2", 68): "B",
}

# Filas del xlsx que son duplicados exactos (mismo pensamiento/competencia/
# clave/retroalimentacion que la fila anterior) y que desfasan la numeración
# real del cuadernillo impreso a partir de ese punto. Confirmado para grado 9
# sesión 2 con 3 evidencias independientes: (1) el cuadernillo impreso
# (Cuadernillo_N2_Grado9_CESUM.pdf) solo llega hasta la pregunta 65, con
# ARTÍSTICA=6 preguntas (28-33, no 28-34), RELIGION=6 (54-59, no 55-61) y
# TECNOLOGÍA=6 (60-65, no 62-68); (2) las filas 33/34, 60/61 y 67/68 de la
# tabla de respuestas son copias exactas fila-a-fila; (3) en el CSV de
# ZipGrade con las respuestas letra por letra, las columnas #66-#68 quedan en
# blanco para ~97-99% de los estudiantes (no hubo pregunta física ahí).
# clave: (grado, sesion) -> set de N° PREGUNTA (numeración original del xlsx)
# a eliminar; todo lo posterior se renumera corriendo la numeración real.
DROP_DUPLICATE_ROWS = {
    ("9", "session2"): {34, 61, 68},
}

# Preguntas que existen como numero impreso en el cuadernillo (y por lo tanto
# como columna en la hoja de respuestas / ZipGrade) pero que NO tienen
# opciones de respuesta reales -- es un error de diagramacion del cuadernillo,
# no un duplicado de fila. Se excluyen del area (no se cuentan ni se
# califican). Confirmado para grado 9 sesion 1: la pregunta 23 es solo el
# texto de contexto de la pregunta 24 ("Laura tiene 16 años..."), sin
# opciones A-D propias. El 79% de los estudiantes (152/193) la dejo en
# blanco; el resto se confundio y marco algo ahi de todas formas.
# clave: (grado, sesion) -> set de preguntas reales a excluir
VOID_QUESTIONS = {
    ("9", "session1"): {23},
}

# Como la fila fantasma de la pregunta 23 tambien desplazo en 1 la CLAVE
# (no solo el conteo) de las preguntas 24-27 en la tabla original -- la fila
# vieja "23"=D en realidad es la clave real de la pregunta 24, la vieja
# "24"=A es la clave real de la 25, etc. -- se sobreescriben directamente
# con la clave ya realineada (la vieja fila 27=C, que sobra, se descarta).
# Confirmado con el consenso de respuestas de los estudiantes: 81% eligio D
# en la 24, 89% eligio A en la 25, 87% eligio C en la 26, 89% eligio D en la 27.
# clave: (grado, sesion) -> {pregunta real: clave correcta}
CLAVE_OVERRIDES = {
    ("9", "session1"): {24: "D", 25: "A", 26: "C", 27: "D"},
}


def norm_clave(v):
    if v is None:
        return None
    s = str(v).strip().upper()
    return s or None


def load_sheet_rows(xlsx_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    ws = wb["Hoja1"]
    # fila 6 es el encabezado ('AREA','N° PREGUNTA',...); los datos empiezan en la 7
    return list(ws.iter_rows(min_row=7, max_row=ws.max_row, values_only=True))


def split_sessions(rows):
    """Devuelve (rows_session1, rows_session2) a partir de las filas crudas."""
    # Caso grado 10/11: marcador explícito 'CUADERNILLO 2'
    for i, row in enumerate(rows):
        area = row[0]
        if area == "CUADERNILLO 2":
            return rows[:i], rows[i + 1 :]

    # Caso grado 6-9: sesión 2 empieza cuando el área 'GESTION' aparece
    for i, row in enumerate(rows):
        if row[0] == "GESTION":
            return rows[:i], rows[i:]

    raise ValueError("No se encontró el límite de sesión 2 (ni CUADERNILLO 2 ni GESTION)")


def build_areas_and_key(rows, grade, session_name):
    areas = []
    key = {}
    current_area = None
    current_start = None
    prev_num = None
    drop_set = DROP_DUPLICATE_ROWS.get((grade, session_name), set())
    dropped = []
    real_num = 0  # numeración real, corrida tras eliminar duplicados

    def close_area(end_num):
        if current_area is not None:
            areas.append(
                {"area": current_area, "start": current_start, "end": end_num}
            )

    for row in rows:
        area, orig_num, _pensamiento, _competencia, clave, _retro = row[:6]
        if orig_num is None:
            continue
        orig_num = int(orig_num)

        if orig_num in drop_set:
            dropped.append(orig_num)
            continue

        real_num += 1
        if area:
            close_area(prev_num)
            current_area = str(area).strip()
            current_start = real_num
        c = norm_clave(clave)
        fix = MANUAL_FIXES.get((grade, session_name, real_num))
        if fix:
            c = fix
        if c is None:
            print(f"AVISO: grado {grade} {session_name} pregunta {real_num} sin clave")
        key[str(real_num)] = c
        prev_num = real_num
    close_area(prev_num)

    if drop_set:
        faltantes = drop_set - set(dropped)
        if faltantes:
            raise ValueError(
                f"grado {grade} {session_name}: no se encontraron las filas a "
                f"eliminar {faltantes} (¿cambió el archivo fuente?)"
            )
        print(f"  grado {grade} {session_name}: eliminadas {len(dropped)} filas duplicadas "
              f"({dropped}), renumerado a {real_num} preguntas reales")

    overrides = CLAVE_OVERRIDES.get((grade, session_name), {})
    for q, clave in overrides.items():
        key[str(q)] = clave
    if overrides:
        print(f"  grado {grade} {session_name}: clave realineada manualmente para preguntas {sorted(overrides)}")

    void = VOID_QUESTIONS.get((grade, session_name), set())
    for q in void:
        key.pop(str(q), None)
    if void:
        print(f"  grado {grade} {session_name}: preguntas anuladas (sin opciones reales) {sorted(void)}")

    return areas, key


def main():
    with zipfile.ZipFile(SOURCE_ZIP) as zf:
        result = {}
        for grade, path in GRADE_FILES.items():
            xlsx_bytes = zf.read(path)
            rows = load_sheet_rows(xlsx_bytes)
            s1_rows, s2_rows = split_sessions(rows)

            s1_areas, s1_key = build_areas_and_key(s1_rows, grade, "session1")
            s2_areas, s2_key = build_areas_and_key(s2_rows, grade, "session2")

            # "total" = ultima columna fisica a leer del CSV de ZipGrade (el
            # maximo "end" de area), NO len(key) -- si una pregunta en medio
            # del rango quedo anulada (VOID_QUESTIONS), len(key) seria menor
            # que el numero de columnas reales y truncaria la lectura de las
            # preguntas siguientes.
            s1_total = max((a["end"] for a in s1_areas), default=0)
            s2_total = max((a["end"] for a in s2_areas), default=0)

            result[grade] = {
                "session1": {
                    "total": s1_total,
                    "areas": s1_areas,
                    "key": s1_key,
                },
                "session2": {
                    "total": s2_total,
                    "areas": s2_areas,
                    "key": s2_key,
                },
            }
            print(
                f"Grado {grade}: sesion1={s1_total}q ({len(s1_key)} con clave, "
                f"{len(s1_areas)} areas), sesion2={s2_total}q ({len(s2_key)} con clave, "
                f"{len(s2_areas)} areas)"
            )

    OUT_PATH.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\nEscrito: {OUT_PATH}")


if __name__ == "__main__":
    main()
